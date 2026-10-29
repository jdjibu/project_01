"""Phase 4-5: build the Makita three-statement + DCF + comps + sensitivity model with live Excel formulas.

    python src/build_model.py            # builds model/Makita_6586_model.xlsx (Base), recalculates it,
                                         # runs Bull/Bear, writes data/clean/model_outputs.json

Conventions: blue = hard-coded input (historical fact or assumption), black = formula, green = link to another
sheet. Historical columns C:G = FY2022-FY2026 (actual), forecast H:L = FY2027E-FY2031E. JPY millions.
Every forecast cell is a formula driven by the Assumptions sheet; the scenario switch is Assumptions!D4.
"""
import json
import shutil
from datetime import date
from pathlib import Path

import openpyxl
import pandas as pd
import xlsxwriter
from xlsxwriter.utility import xl_rowcol_to_cell

from recalc import recalc

ROOT = Path(__file__).resolve().parents[1]
CLEAN, MODEL = ROOT / "data/clean", ROOT / "model"
OUT = MODEL / "Makita_6586_model.xlsx"

HIST = [2022, 2023, 2024, 2025, 2026]
FCST = [2027, 2028, 2029, 2030, 2031]
YEARS = HIST + FCST
COL = {y: 2 + i for i, y in enumerate(YEARS)}  # C..L
L = {y: chr(ord("A") + COL[y]) for y in YEARS}  # column letters
NOTE_COL = 12  # M
F0, F1 = L[FCST[0]], L[FCST[-1]]


# ----------------------------------------------------------------------------------------------- data
def load_hist():
    IS = pd.read_csv(CLEAN / "makita_is.csv", index_col=0)
    BS = pd.read_csv(CLEAN / "makita_bs.csv", index_col=0).fillna(0)
    CF = pd.read_csv(CLEAN / "makita_cf.csv", index_col=0).fillna(0)
    sup = pd.read_csv(CLEAN / "makita_supplementary_wide.csv", index_col=0)
    geo = pd.read_csv(CLEAN / "makita_geo_revenue.csv")
    geo = geo[geo.release_fy == geo.fy].pivot(index="region", columns="fy", values="value")
    snap = pd.read_csv(CLEAN / "market_snapshot.csv").set_index("item")["value"]
    q = pd.read_csv(CLEAN / "makita_quarterly.csv")
    bsq = q[(q.period == "1Q2027") & (q.statement == "BS")].set_index("item")["value"]
    ltm = pd.read_csv(CLEAN / "ltm_makita.csv").set_index("item")["LTM_Jun2026"]
    dam = pd.read_csv(CLEAN / "damodaran_inputs.csv").set_index("item")["value"]
    crp = pd.read_csv(CLEAN / "crp_weighted.csv").set_index("makita_region")["crp"]
    beta = pd.read_csv(CLEAN / "beta_regressions.csv")
    rf = pd.read_csv(CLEAN / "rf_jgb10y.csv").iloc[-1]
    peers = pd.read_csv(CLEAN / "peers_financials.csv")
    hm = pd.read_csv(CLEAN / "makita_hist_multiples.csv")
    fxd = pd.read_csv(CLEAN / "fx_daily.csv", index_col=0)
    return dict(IS=IS, BS=BS, CF=CF, sup=sup, geo=geo, snap=snap, bsq=bsq, ltm=ltm, dam=dam, crp=crp,
                beta=beta, rf=rf, peers=peers, hm=hm, fxd=fxd)


# ----------------------------------------------------------------------------------------------- helpers
DEFERRED = []  # formulas are written after every sheet's rows are known (sheets reference each other)


def flush():
    while DEFERRED:
        ws, row, col, fn, y, style = DEFERRED.pop(0)
        ws.write_formula(row, col, fn(y), style)


class Sheet:
    """Tracks named rows so formulas can reference lines by key."""

    def __init__(self, wb, name, fmt):
        self.ws, self.name, self.f, self.rows, self.r = wb.add_worksheet(name), name, fmt, {}, 0

    def ref(self, key, year=None, col=None, absolute=False, sheet=False):
        row = self.rows[key]
        c = COL[year] if year is not None else col
        cell = xl_rowcol_to_cell(row, c, row_abs=absolute, col_abs=absolute)
        return f"'{self.name}'!{cell}" if sheet else cell

    def title(self, text, sub=None):
        self.ws.write(0, 0, text, self.f["title"])
        if sub:
            self.ws.write(1, 0, sub, self.f["sub"])
        self.r = 3

    def header(self, label="JPY m, FY ending March", years=True):
        self.ws.write(self.r, 0, label, self.f["hdr"])
        self.ws.write(self.r, 1, "", self.f["hdr"])
        if years:
            for y in YEARS:
                self.ws.write(self.r, COL[y], f"FY{y}{'A' if y in HIST else 'E'}", self.f["hdr_r"])
        self.ws.write(self.r, NOTE_COL, "Source / note", self.f["hdr"])
        self.r += 1

    def section(self, text):
        self.r += 1
        self.ws.write(self.r, 0, text, self.f["sec"])
        self.r += 1

    def line(self, key, label, hist=None, fcst=None, fmt="num", note="", bold=False, unit="", hist_formula=None):
        """hist: dict year->value (input), or hist_formula: fn(year)->formula; fcst: fn(year)->formula."""
        row = self.r
        self.rows[key] = row
        style = self.f[("b_" if bold else "") + fmt]
        self.ws.write(row, 0, label, self.f["bold"] if bold else self.f["label"])
        self.ws.write(row, 1, unit, self.f["unit"])
        for y in HIST:
            if hist_formula is not None:
                DEFERRED.append((self.ws, row, COL[y], hist_formula, y, style))
            elif hist is not None and y in hist and pd.notna(hist[y]):
                self.ws.write_number(row, COL[y], float(hist[y]), self.f[("b_" if bold else "") + fmt + "_in"])
        if fcst is not None:
            for y in FCST:
                DEFERRED.append((self.ws, row, COL[y], fcst, y, style))
        if note:
            self.ws.write(row, NOTE_COL, note, self.f["note"])
        self.r += 1
        return row


def prev(y):
    return y - 1


def formats(wb):
    base = {"font_name": "Arial", "font_size": 9}
    f = {}
    f["title"] = wb.add_format({**base, "bold": True, "font_size": 14})
    f["sub"] = wb.add_format({**base, "italic": True, "font_color": "#5C5C5C"})
    f["hdr"] = wb.add_format({**base, "bold": True, "bg_color": "#1A1A1A", "font_color": "#FFFFFF"})
    f["hdr_r"] = wb.add_format({**base, "bold": True, "bg_color": "#1A1A1A", "font_color": "#FFFFFF", "align": "right"})
    f["sec"] = wb.add_format({**base, "bold": True, "font_color": "#A8792A", "bottom": 1, "bottom_color": "#A8792A"})
    f["label"] = wb.add_format(base)
    f["bold"] = wb.add_format({**base, "bold": True})
    f["unit"] = wb.add_format({**base, "font_color": "#8C8C8C"})
    f["note"] = wb.add_format({**base, "font_color": "#5C5C5C", "italic": True})
    numfmts = {"num": "#,##0;(#,##0);-", "pct": "0.0%;(0.0%);-", "x": '0.0"x"', "dec": "#,##0.00", "days": "0",
               "date": "dd-mmm-yyyy", "yen": "#,##0", "dec3": "0.000"}
    for k, nf in numfmts.items():
        f[k] = wb.add_format({**base, "num_format": nf})
        f[k + "_in"] = wb.add_format({**base, "num_format": nf, "font_color": "#1F4E9E"})
        f["b_" + k] = wb.add_format({**base, "num_format": nf, "bold": True, "top": 1})
        f["b_" + k + "_in"] = wb.add_format({**base, "num_format": nf, "bold": True, "top": 1, "font_color": "#1F4E9E"})
        f[k + "_link"] = wb.add_format({**base, "num_format": nf, "font_color": "#2E7D32"})
        f[k + "_key"] = wb.add_format({**base, "num_format": nf, "bold": True, "bg_color": "#F3EBDD", "border": 1})
    f["switch"] = wb.add_format({**base, "bold": True, "bg_color": "#FFF2CC", "border": 2, "font_color": "#1F4E9E",
                                 "align": "center"})
    f["ok"] = wb.add_format({**base, "bold": True, "font_color": "#2E7D32"})
    f["text_in"] = wb.add_format({**base, "font_color": "#1F4E9E"})
    f["wrap"] = wb.add_format({**base, "text_wrap": True, "valign": "top"})
    return f


# ----------------------------------------------------------------------------------------------- assumptions
# Time-series drivers: (key, label, unit, fmt, {scenario: [FY27..FY31]}, justification)
TS = [
    ("lcg_eur", "Local-currency growth: Europe", "%", "pct",
     {"Base": [0.000, 0.025, 0.030, 0.030, 0.030], "Bull": [0.020, 0.045, 0.050, 0.050, 0.040],
      "Bear": [-0.030, -0.005, 0.015, 0.020, 0.020]},
     "Implied LC growth FY2026 -1.6%, 1Q FY2027 about -2% (Drivers sheet). Base: flat FY27, then cordless/OPE-led recovery to ~3%"),
    ("lcg_jpn", "Local-currency growth: Japan", "%", "pct",
     {"Base": [0.025, 0.020, 0.020, 0.020, 0.015], "Bull": [0.040, 0.035, 0.030, 0.030, 0.025],
      "Bear": [0.000, 0.000, 0.010, 0.010, 0.010]},
     "FY2026 +3.9%, 1Q FY2027 +4.1%; housing starts falling. Base fades to ~2%"),
    ("lcg_na", "Local-currency growth: North America", "%", "pct",
     {"Base": [0.030, 0.030, 0.030, 0.030, 0.030], "Bull": [0.060, 0.060, 0.060, 0.050, 0.040],
      "Bear": [-0.030, 0.000, 0.010, 0.020, 0.020]},
     "FY2026 LC -5%, 1Q FY2027 LC about +7% after tariff disruption. Base 3%: share stable vs TTI/DeWalt"),
    ("lcg_row", "Local-currency growth: Rest of world", "%", "pct",
     {"Base": [0.020, 0.035, 0.040, 0.040, 0.035], "Bull": [0.040, 0.050, 0.055, 0.055, 0.050],
      "Bear": [-0.020, 0.010, 0.020, 0.025, 0.025]},
     "Asia, Latin America, Oceania, Middle East & Africa; USD used as FX proxy. FY2026 implied LC +5%, 1Q FY2027 ~0%"),
    ("fx_usd", "USD/JPY (fiscal-year average)", "JPY", "dec",
     {"Base": [158.26, 157.82, 157.82, 157.82, 157.82], "Bull": [158.26, 165.71, 165.71, 165.71, 165.71],
      "Bear": [154.31, 142.04, 142.04, 142.04, 142.04]},
     "Base: 1Q actual 159.57 blended with spot 157.82 (29 Oct 2026) for 3 quarters, then spot held (random walk). Bull: yen -5%. Bear: yen +10%"),
    ("fx_eur", "EUR/JPY (fiscal-year average)", "JPY", "dec",
     {"Base": [178.38, 176.03, 176.03, 176.03, 176.03], "Bull": [178.38, 184.83, 184.83, 184.83, 184.83],
      "Bear": [174.08, 158.43, 158.43, 158.43, 158.43]},
     "Base: 1Q actual 185.41 blended with spot 176.03 for 3 quarters, then spot held. Bull: yen -5%. Bear: yen +10%"),
    ("gm", "Gross margin", "%", "pct",
     {"Base": [0.370, 0.365, 0.365, 0.365, 0.365], "Bull": [0.375, 0.375, 0.380, 0.380, 0.380],
      "Bear": [0.355, 0.340, 0.350, 0.350, 0.350]},
     "FY2026 36.5%; 1Q FY2027 39.3% helped by one-off US tariff refund (non-recurring). Bear: FY2024-style inventory correction (30.2% then) but milder"),
    ("sga", "SG&A, % of revenue", "%", "pct",
     {"Base": [0.233, 0.230, 0.228, 0.226, 0.225], "Bull": [0.230, 0.225, 0.220, 0.218, 0.215],
      "Bear": [0.238, 0.238, 0.235, 0.235, 0.235]},
     "FY2026 23.1%, 1Q FY2027 24.5% (sales headcount build under Vision FY2030). Base: modest operating leverage"),
    ("dso", "Receivable days (on revenue)", "days", "days",
     {"Base": [57] * 5, "Bull": [55] * 5, "Bear": [60] * 5}, "FY2018-26 range 51-60; FY2026 57"),
    ("dio", "Inventory days (on cost of sales)", "days", "days",
     {"Base": [275, 268, 262, 258, 255], "Bull": [270, 258, 250, 248, 245], "Bear": [290, 290, 285, 285, 285]},
     "FY2026 278, FY2018-26 average ex FY2022-23 = 249. Base: gradual normalisation to 255"),
    ("dpo", "Payable days (on cost of sales)", "days", "days",
     {"Base": [41] * 5, "Bull": [42] * 5, "Bear": [39] * 5}, "FY2025-26 41"),
    ("oca", "Other current assets, % of revenue", "%", "pct",
     {"Base": [0.025] * 5, "Bull": [0.025] * 5, "Bear": [0.025] * 5}, "FY2026 2.5%"),
    ("ool", "Other operating liabilities, % of revenue", "%", "pct",
     {"Base": [0.098] * 5, "Bull": [0.098] * 5, "Bear": [0.098] * 5},
     "Tax payable + provisions + other current liabilities; FY2026 9.8%"),
    ("capex", "Capex, % of revenue", "%", "pct",
     {"Base": [0.036, 0.034, 0.032, 0.032, 0.032], "Bull": [0.036, 0.035, 0.034, 0.034, 0.034],
      "Bear": [0.033, 0.028, 0.028, 0.028, 0.028]},
     "Company guides JPY 30bn for FY2027 (about 3.6% of base revenue); FY2018-26 average 5.0%, FY2024-26 2.5%"),
    ("buyback", "Share buybacks (JPY m)", "JPY m", "num",
     {"Base": [0] * 5, "Bull": [20000] * 5, "Bear": [0] * 5},
     "No programme announced beyond 1Q FY2027 purchases; base assumes none (does not change DCF value per share)"),
]

# single-value assumptions: (key, label, fmt, base, bull, bear, justification)
SV = [
    ("tax", "Effective tax rate", "pct", 0.270, 0.265, 0.280, "FY2025 27.0%, FY2026 26.5%; Japan statutory about 30.6%, lower overseas rates"),
    ("dep_rate", "Depreciation ex right-of-use, % of opening PP&E", "pct", 0.093, 0.093, 0.093,
     "FY2026: 25,205 / 266,609 = 9.5%; company guides D&A JPY 25.0bn for FY2027"),
    ("rou", "Right-of-use depreciation = lease payments, % of revenue", "pct", 0.0068, 0.0068, 0.0068,
     "FY2026: CF D&A 30,465 less D&A ex-ROU 25,205 = 5,260 (0.68%); lease repayments 5,151"),
    ("cash_yield", "Interest yield on opening cash and deposits", "pct", 0.018, 0.020, 0.015,
     "FY2026 interest received 5,936 on about JPY 292bn (2.0%); lower overseas rates assumed"),
    ("finexp", "Financial expenses (JPY m per year)", "num", -1200, -1200, -1200, "Lease interest 981 (FY2026 note) + other"),
    ("nci", "NCI share of profit", "pct", 0.001, 0.001, 0.001, "FY2025-26 about 0%"),
    ("payout", "Dividend payout ratio", "pct", 0.50, 0.50, 0.50, "Policy: consolidated payout of 50% or more (FY2026 release)"),
    ("min_cash", "Minimum operating cash (JPY m)", "num", 100000, 100000, 100000, "Revolver draws if cash falls below this"),
    ("tg", "Terminal growth (Gordon)", "pct", 0.020, 0.025, 0.015,
     "Long-run nominal growth: ~2% inflation in core markets, modest real growth offset by share loss; below rf 2.87%"),
    ("ronic", "Return on new invested capital in terminal period", "pct", 0.12, 0.15, 0.10,
     "Value-driver TV: reinvestment = g / RONIC. Makita ROIC FY2018-26: median 12.2%, range 1.8-18.4% (ratios.csv)"),
    ("exit", "Exit EV/EBITDA multiple (FY2031)", "x", 7.5, 8.5, 6.5,
     "Set within ~0.5x of each scenario's Gordon-implied multiple so the two TV methods agree (exit-implied g <= rf). "
     "Current LTM 7.8x; Makita FY2019-26 median 9.8x would assume a re-rating; peer median ~13x is TTI/Toro-heavy"),
]


def build(d, scenario: int, path: Path, overrides=None):
    """overrides: {driver key: {scenario: [5 values]}} or {single key: (base, bull, bear)} - used by src/risks.py"""
    overrides = overrides or {}
    wb = xlsxwriter.Workbook(str(path))
    f = formats(wb)
    IS_, BS_, CF_, sup, geo = d["IS"], d["BS"], d["CF"], d["sup"], d["geo"]
    snap, bsq, ltm, dam, crp, rf = d["snap"], d["bsq"], d["ltm"], d["dam"], d["crp"], d["rf"]
    h = lambda df, item: {y: df.at[item, f"FY{y}"] for y in HIST if item in df.index}

    cover = wb.add_worksheet("Cover")
    A = Sheet(wb, "Assumptions", f)
    D = Sheet(wb, "Drivers", f)
    I = Sheet(wb, "IS", f)
    B = Sheet(wb, "BS", f)
    C = Sheet(wb, "CF", f)
    V = Sheet(wb, "DCF", f)
    P = Sheet(wb, "Comps", f)
    S = Sheet(wb, "Sensitivity", f)
    FF = Sheet(wb, "Football", f)
    K = Sheet(wb, "Checks", f)
    for sh in [A, D, I, B, C, V, P, S, FF, K]:
        sh.ws.set_column(0, 0, 46)
        sh.ws.set_column(1, 1, 8)
        sh.ws.set_column(2, 11, 11)
        sh.ws.set_column(NOTE_COL, NOTE_COL, 90)
        sh.ws.freeze_panes(4, 2)

    # ---------------- Assumptions
    A.title("Assumptions", "Every forecast driver lives here. Change the scenario in D4. Blue = input.")
    A.ws.write(3, 0, "Scenario switch: 1 = Base, 2 = Bull, 3 = Bear", f["bold"])
    A.ws.write_number(3, 3, scenario, f["switch"])
    A.ws.write_formula(3, 4, '=CHOOSE($D$4,"Base","Bull","Bear")', f["ok"])
    A.ws.data_validation(3, 3, 3, 3, {"validate": "list", "source": [1, 2, 3]})
    wb.define_name("Scenario", "=Assumptions!$D$4")
    A.r = 5
    A.header("Time-series drivers (live row = CHOOSE on scenario)")
    for key, label, unit, fmt, vals, note in TS:
        vals = {**vals, **overrides.get(key, {})}
        live = A.r
        A.rows[key] = live
        A.ws.write(live, 0, label + " (live)", f["bold"])
        A.ws.write(live, 1, unit, f["unit"])
        for i, y in enumerate(FCST):
            c = COL[y]
            A.ws.write_formula(live, c, f"=CHOOSE(Scenario,{xl_rowcol_to_cell(live + 1, c)},{xl_rowcol_to_cell(live + 2, c)},"
                                        f"{xl_rowcol_to_cell(live + 3, c)})", f["b_" + fmt])
        A.ws.write(live, NOTE_COL, note, f["note"])
        for j, scen in enumerate(["Base", "Bull", "Bear"]):
            A.ws.write(live + 1 + j, 0, f"   {scen}", f["label"])
            for i, y in enumerate(FCST):
                A.ws.write_number(live + 1 + j, COL[y], vals[scen][i], f[fmt + "_in"])
        A.r += 4
    A.section("Single-value assumptions        (D = Base, E = Bull, F = Bear, G = live)")
    A.ws.write_row(A.r, 3, ["Base", "Bull", "Bear", "Live"], f["hdr_r"])
    A.r += 1
    for key, label, fmt, b, u, e, note in SV:
        b, u, e = overrides.get(key, (b, u, e))
        row = A.r
        A.rows[key] = row
        A.ws.write(row, 0, label, f["label"])
        for j, v in enumerate([b, u, e]):
            A.ws.write_number(row, 3 + j, v, f[fmt + "_in"])
        A.ws.write_formula(row, 6, f"=CHOOSE(Scenario,D{row + 1},E{row + 1},F{row + 1})", f["b_" + fmt])
        A.ws.write(row, NOTE_COL, note, f["note"])
        A.r += 1
    sv = lambda key: f"Assumptions!$G${A.rows[key] + 1}"
    ts = lambda key, y: f"Assumptions!{L[y]}{A.rows[key] + 1}"

    # market / WACC inputs (same in all scenarios; facts with sources)
    A.section("Market and cost-of-capital inputs (facts, sourced)")
    MK = {}

    def mk(key, label, value, fmt, note):
        row = A.r
        MK[key] = f"Assumptions!$D${row + 1}"
        A.ws.write(row, 0, label, f["label"])
        if isinstance(value, str) and value.startswith("="):
            A.ws.write_formula(row, 3, value, f[fmt])
        elif fmt == "date":
            A.ws.write_datetime(row, 3, value, f["date_in"])
        else:
            A.ws.write_number(row, 3, float(value), f[fmt + "_in"])
        A.ws.write(row, NOTE_COL, note, f["note"])
        A.r += 1

    beta_reg = d["beta"].query("index == 'TOPIX ETF (1306.T)' and frequency == 'weekly' and years == 5").iloc[0]
    mk("val_date", "Valuation date", date(2026, 10, 29), "date", "Tokyo close 29 Oct 2026")
    mk("bs_date", "Balance-sheet date for equity bridge", date(2026, 6, 30), "date", "Latest filed balance sheet (1Q FY2027)")
    mk("price", "Share price (JPY)", snap["price_close_jpy"], "yen", "yfinance 6586.T close, 29 Oct 2026")
    mk("shares", "Shares outstanding (m)", float(snap["shares_outstanding_30jun2026"]) / 1e6, "dec",
       "1Q FY2027 release: 280,017,520 issued less 22,362,009 treasury")
    mk("rf", "Risk-free rate: 10Y JGB", float(rf["10Y"]) / 100, "pct", f"Ministry of Finance Japan, {rf['Date']}")
    mk("erp", "Mature-market equity risk premium", dam["Mature market ERP (= Japan total ERP - Japan CRP)"], "pct",
       "Damodaran ctryprem.xlsx (implied S&P 500 ERP, start 2026)")
    mk("crp", "Country risk premium, revenue-weighted", crp["Revenue-weighted CRP"], "pct",
       "Damodaran regional CRPs weighted by Makita FY2026 revenue by customer region (crp_weighted.csv)")
    mk("bu", "Unlevered beta: Damodaran Japan Machinery", dam["Japan Machinery: unlevered beta"], "dec3",
       "betaJapan.xls, 221 firms. Cross-check: 5y weekly regression vs TOPIX ETF = %.2f (s.e. %.2f)" % (beta_reg.raw_beta, beta_reg.std_err))
    mk("debt_mv", "Debt incl. leases at 30 Jun 2026 (JPY m)", bsq["borrowings_cur"] + bsq["other_fin_liab_cur"] + bsq["other_fin_liab_nc"],
       "num", "Borrowings 1,604 + other financial liabilities (mainly leases) 20,378")
    mk("mcap", "Market capitalisation (JPY m)", f"={MK['price']}*{MK['shares']}", "num", "Price x shares")
    mk("de", "Debt / equity (market)", f"={MK['debt_mv']}/{MK['mcap']}", "pct", "")
    mk("bl", "Relevered beta", f"={MK['bu']}*(1+(1-{sv('tax')})*{MK['de']})", "dec3", "Hamada: bu x (1 + (1 - t) D/E)")
    mk("size", "Size premium", 0, "pct", "None: market cap about JPY 1.35tn (large cap)")
    mk("ke", "Cost of equity", f"={MK['rf']}+{MK['bl']}*({MK['erp']}+{MK['crp']})+{MK['size']}", "pct", "CAPM: rf + beta x (ERP + CRP)")
    mk("spread", "Pre-tax credit spread", 0.0075, "pct", "Net cash, interest cover about 90x: high-grade spread assumed")
    mk("kd", "After-tax cost of debt", f"=({MK['rf']}+{MK['spread']})*(1-{sv('tax')})", "pct", "")
    mk("wd", "Weight of debt", f"={MK['debt_mv']}/({MK['debt_mv']}+{MK['mcap']})", "pct", "Market-value weights")
    mk("we", "Weight of equity", f"=1-{MK['wd']}", "pct", "")
    mk("wacc", "WACC", f"={MK['we']}*{MK['ke']}+{MK['wd']}*{MK['kd']}", "pct_key", "")
    mk("cash_q", "Cash at 30 Jun 2026", bsq["cash"], "num", "1Q FY2027 balance sheet")
    mk("dep_q", "Other current financial assets (mainly time deposits)", bsq["other_fin_assets_cur"], "num", "1Q FY2027 balance sheet")
    mk("inv_q", "Other non-current financial assets (investments)", bsq["other_fin_assets_nc"], "num",
       "1Q FY2027 balance sheet; non-operating, at book")
    mk("borr_q", "Borrowings at 30 Jun 2026", bsq["borrowings_cur"], "num", "1Q FY2027 balance sheet")
    mk("nci_q", "Non-controlling interests (book)", bsq["nci"], "num", "1Q FY2027 balance sheet")
    mk("lease_fy", "Lease liabilities (FY2026 note 27)", 19104, "num", "Audited FS FY3/2026; held flat in forecast")
    mk("interim", "FY2027 interim DPS already ex-dividend (JPY)", 79, "yen", "Declared; record date 30 Sep 2026, before valuation date")

    # ---------------- Drivers (revenue build)
    D.title("Drivers: revenue build by customer region", "Revenue_t = Revenue_t-1 x (1 + local-currency growth) x (FX_t / FX_t-1)")
    D.header()
    D.section("FX (fiscal-year average)")
    D.line("usd", "USD/JPY", hist=h(sup, "fx_usdjpy"), fcst=lambda y: f"={ts('fx_usd', y)}", fmt="dec",
           note="Historical: Makita releases; forecast: Assumptions")
    D.line("eur", "EUR/JPY", hist=h(sup, "fx_eurjpy"), fcst=lambda y: f"={ts('fx_eur', y)}", fmt="dec")
    regions = [("eu", "Europe", "eur", "lcg_eur", ["Europe"]), ("jp", "Japan", None, "lcg_jpn", ["Japan"]),
               ("na", "North America", "usd", "lcg_na", ["North America"]),
               ("rw", "Rest of world", "usd", "lcg_row", ["Asia", "Central & South America", "Oceania", "Middle East & Africa"])]
    for key, name, fx, g, members in regions:
        D.section(name + (f" (FX proxy: {fx.upper()})" if fx else " (JPY)"))
        hist = {y: sum(geo.at[m, y] for m in members) for y in HIST}
        fxf = (lambda y, fx=fx: f"*({D.ref(fx, y)}/{D.ref(fx, prev(y))})") if fx else (lambda y: "")
        D.line(key, f"Revenue: {name}", hist=hist, fcst=lambda y, key=key, g=g, fxf=fxf:
               f"={D.ref(key, prev(y))}*(1+{ts(g, y)}){fxf(y)}", note="Makita revenue by customer location")
        D.line(key + "_g", "  Growth in JPY", hist_formula=lambda y, key=key: f"=IFERROR({D.ref(key, y)}/{D.ref(key, prev(y))}-1,\"\")" if y > HIST[0] else '=""',
               fcst=lambda y, key=key: f"={D.ref(key, y)}/{D.ref(key, prev(y))}-1", fmt="pct")
        D.line(key + "_lc", "  Implied local-currency growth", fmt="pct",
               hist_formula=(lambda y, key=key, fx=fx: (f"=({D.ref(key, y)}/{D.ref(key, prev(y))})/({D.ref(fx, y)}/{D.ref(fx, prev(y))})-1"
                                                          if fx else f"={D.ref(key, y)}/{D.ref(key, prev(y))}-1") if y > HIST[0] else '=""'),
               fcst=lambda y, g=g: f"={ts(g, y)}", note="Historical: JPY growth stripped of FX; forecast: assumption")
    D.section("Total")
    D.line("rev", "Total revenue", hist_formula=lambda y: "=" + "+".join(D.ref(k, y) for k, *_ in regions),
           fcst=lambda y: "=" + "+".join(D.ref(k, y) for k, *_ in regions), bold=True)
    D.line("rev_g", "  Growth", hist_formula=lambda y: f"={D.ref('rev', y)}/{D.ref('rev', prev(y))}-1" if y > HIST[0] else '=""',
           fcst=lambda y: f"={D.ref('rev', y)}/{D.ref('rev', prev(y))}-1", fmt="pct")
    D.line("fx_contrib", "  of which FX translation (approx.)", fmt="pct",
           fcst=lambda y: f"=({D.ref('rev', y)}-SUMPRODUCT({D.ref('eu', prev(y))}*(1+{ts('lcg_eur', y)})+{D.ref('jp', prev(y))}*(1+{ts('lcg_jpn', y)})"
                          f"+{D.ref('na', prev(y))}*(1+{ts('lcg_na', y)})+{D.ref('rw', prev(y))}*(1+{ts('lcg_row', y)})))/{D.ref('rev', prev(y))}",
           note="Revenue growth from FX alone, holding local-currency volumes")

    # ---------------- Income statement
    I.title("Income statement", "IFRS. Historical from Makita earnings releases (data/clean/makita_is.csv)")
    I.header()
    I.line("rev", "Revenue", hist=h(IS_, "revenue"), fcst=lambda y: f"={D.ref('rev', y, sheet=True)}", bold=True,
           note="Forecast: Drivers sheet regional build")
    I.line("rev_g", "  Growth", hist_formula=lambda y: f"=IFERROR({I.ref('rev', y)}/{I.ref('rev', prev(y))}-1,\"\")" if y > HIST[0] else '=""',
           fcst=lambda y: f"={I.ref('rev', y)}/{I.ref('rev', prev(y))}-1", fmt="pct")
    I.line("cogs", "Cost of sales", hist=h(IS_, "cost_of_sales"), fcst=lambda y: f"=-{I.ref('rev', y)}*(1-{ts('gm', y)})")
    I.line("gp", "Gross profit", hist_formula=lambda y: f"={I.ref('rev', y)}+{I.ref('cogs', y)}",
           fcst=lambda y: f"={I.ref('rev', y)}+{I.ref('cogs', y)}", bold=True)
    I.line("gm", "  Gross margin", hist_formula=lambda y: f"={I.ref('gp', y)}/{I.ref('rev', y)}",
           fcst=lambda y: f"={I.ref('gp', y)}/{I.ref('rev', y)}", fmt="pct")
    I.line("sga", "SG&A and other, net", hist=h(IS_, "sga_net"), fcst=lambda y: f"=-{I.ref('rev', y)}*{ts('sga', y)}")
    I.line("op", "Operating profit", hist_formula=lambda y: f"={I.ref('gp', y)}+{I.ref('sga', y)}",
           fcst=lambda y: f"={I.ref('gp', y)}+{I.ref('sga', y)}", bold=True)
    I.line("opm", "  Operating margin", hist_formula=lambda y: f"={I.ref('op', y)}/{I.ref('rev', y)}",
           fcst=lambda y: f"={I.ref('op', y)}/{I.ref('rev', y)}", fmt="pct")
    I.line("fin_inc", "Financial income", hist=h(IS_, "financial_income"),
           fcst=lambda y: f"={sv('cash_yield')}*({B.ref('cash', prev(y), sheet=True)}+{B.ref('dep', prev(y), sheet=True)})",
           note="Forecast: yield x opening cash and deposits (no circularity)")
    I.line("fin_exp", "Financial expenses", hist=h(IS_, "financial_expenses"), fcst=lambda y: f"={sv('finexp')}")
    I.line("pbt", "Profit before tax", hist_formula=lambda y: f"={I.ref('op', y)}+{I.ref('fin_inc', y)}+{I.ref('fin_exp', y)}",
           fcst=lambda y: f"={I.ref('op', y)}+{I.ref('fin_inc', y)}+{I.ref('fin_exp', y)}", bold=True)
    I.line("tax", "Income tax", hist=h(IS_, "income_tax"), fcst=lambda y: f"=-{I.ref('pbt', y)}*{sv('tax')}")
    I.line("etr", "  Effective tax rate", hist_formula=lambda y: f"=-{I.ref('tax', y)}/{I.ref('pbt', y)}",
           fcst=lambda y: f"=-{I.ref('tax', y)}/{I.ref('pbt', y)}", fmt="pct")
    I.line("profit", "Profit", hist_formula=lambda y: f"={I.ref('pbt', y)}+{I.ref('tax', y)}",
           fcst=lambda y: f"={I.ref('pbt', y)}+{I.ref('tax', y)}", bold=True)
    I.line("ni_nci", "  attributable to non-controlling interests", hist=h(IS_, "ni_nci"),
           fcst=lambda y: f"={I.ref('profit', y)}*{sv('nci')}")
    I.line("ni", "  attributable to owners of the parent", hist_formula=lambda y: f"={I.ref('profit', y)}-{I.ref('ni_nci', y)}",
           fcst=lambda y: f"={I.ref('profit', y)}-{I.ref('ni_nci', y)}", bold=True)
    I.section("Per share")
    I.line("shares", "Average shares (m)", hist={y: sup.at["avg_shares", f"FY{y}"] / 1e6 for y in HIST},
           fcst=lambda y: (f"={MK['shares']}-{ts('buyback', y)}/{MK['price']}/2" if y == FCST[0]
                           else f"={I.ref('shares', prev(y))}-({ts('buyback', prev(y))}+{ts('buyback', y)})/{MK['price']}/2"),
           fmt="dec", note="Forecast: 30 Jun 2026 count less buybacks at today's price (mid-year)")
    I.line("eps", "EPS (JPY)", hist_formula=lambda y: f"={I.ref('ni', y)}/{I.ref('shares', y)}",
           fcst=lambda y: f"={I.ref('ni', y)}/{I.ref('shares', y)}", fmt="dec", bold=True)
    I.line("dps", "DPS (JPY)", hist=h(sup, "dps"), fcst=lambda y: f"={I.ref('eps', y)}*{sv('payout')}", fmt="dec")
    I.section("Memo")
    I.line("da", "D&A total (incl. right-of-use)", hist=h(CF_, "cfo_da"), fcst=lambda y: f"={C.ref('da', y, sheet=True)}")
    I.line("ebitda", "EBITDA (IFRS 16)", hist_formula=lambda y: f"={I.ref('op', y)}+{I.ref('da', y)}",
           fcst=lambda y: f"={I.ref('op', y)}+{I.ref('da', y)}", bold=True)
    I.line("ebitda_m", "  EBITDA margin", hist_formula=lambda y: f"={I.ref('ebitda', y)}/{I.ref('rev', y)}",
           fcst=lambda y: f"={I.ref('ebitda', y)}/{I.ref('rev', y)}", fmt="pct")

    # ---------------- Balance sheet
    bs = lambda *items: {y: sum(BS_.at[i, f"FY{y}"] for i in items if i in BS_.index) for y in HIST}
    B.title("Balance sheet", "Condensed from the IFRS statement of financial position (data/clean/makita_bs.csv)")
    B.header()
    B.section("Assets")
    B.line("cash", "Cash and cash equivalents", hist=bs("cash"), fcst=lambda y: f"={C.ref('cash_end', y, sheet=True)}",
           note="Forecast: from cash-flow statement")
    B.line("dep", "Other current financial assets (mainly time deposits)", hist=bs("other_fin_assets_cur"),
           fcst=lambda y: f"={B.ref('dep', prev(y))}", note="Held flat")
    B.line("rec", "Trade and other receivables", hist=bs("receivables"), fcst=lambda y: f"={I.ref('rev', y, sheet=True)}*{ts('dso', y)}/365")
    B.line("inv", "Inventories", hist=bs("inventories"), fcst=lambda y: f"=-{I.ref('cogs', y, sheet=True)}*{ts('dio', y)}/365")
    B.line("oca", "Other current assets", hist=bs("other_current_assets"), fcst=lambda y: f"={I.ref('rev', y, sheet=True)}*{ts('oca', y)}")
    tca = lambda y: f"=SUM({B.ref('cash', y)}:{B.ref('oca', y)})"
    B.line("tca", "Total current assets", hist_formula=tca, fcst=tca, bold=True)
    B.line("ppe", "Property, plant and equipment (incl. right-of-use)", hist=bs("ppe"),
           fcst=lambda y: f"={B.ref('ppe', prev(y))}+{C.ref('capex_pos', y, sheet=True)}-{C.ref('dep', y, sheet=True)}",
           note="Opening + capex - depreciation ex-ROU (ROU asset held flat: new leases = ROU depreciation)")
    B.line("intang", "Goodwill and intangible assets", hist=bs("intangibles"), fcst=lambda y: f"={B.ref('intang', prev(y))}", note="Held flat")
    B.line("onca", "Other non-current assets (investments, pension, deferred tax)",
           hist=bs("other_fin_assets_nc", "retirement_benefit_assets", "deferred_tax_assets", "other_noncurrent_assets"),
           fcst=lambda y: f"={B.ref('onca', prev(y))}", note="Held flat")
    tnca = lambda y: f"=SUM({B.ref('ppe', y)}:{B.ref('onca', y)})"
    B.line("tnca", "Total non-current assets", hist_formula=tnca, fcst=tnca, bold=True)
    B.line("ta", "Total assets", hist_formula=lambda y: f"={B.ref('tca', y)}+{B.ref('tnca', y)}",
           fcst=lambda y: f"={B.ref('tca', y)}+{B.ref('tnca', y)}", bold=True)
    B.section("Liabilities")
    B.line("pay", "Trade and other payables", hist=bs("payables"), fcst=lambda y: f"=-{I.ref('cogs', y, sheet=True)}*{ts('dpo', y)}/365")
    B.line("borr", "Borrowings", hist=bs("borrowings_cur", "borrowings_nc"),
           fcst=lambda y: f"={B.ref('borr', prev(y))}+{C.ref('d_borr', y, sheet=True)}", note="Base borrowings flat + revolver (CF)")
    B.line("ofl", "Other financial liabilities (mainly leases)", hist=bs("other_fin_liab_cur", "other_fin_liab_nc"),
           fcst=lambda y: f"={B.ref('ofl', prev(y))}", note="Held flat")
    B.line("ool", "Other operating liabilities (tax payable, provisions, other)",
           hist=bs("tax_payable_cur", "provisions_cur", "other_current_liab"),
           fcst=lambda y: f"={I.ref('rev', y, sheet=True)}*{ts('ool', y)}")
    B.line("oncl", "Other non-current liabilities (pension, provisions, deferred tax)",
           hist=bs("retirement_benefit_liab", "provisions_nc", "deferred_tax_liab", "tax_payable_nc", "other_noncurrent_liab"),
           fcst=lambda y: f"={B.ref('oncl', prev(y))}", note="Held flat")
    tl = lambda y: f"=SUM({B.ref('pay', y)}:{B.ref('oncl', y)})"
    B.line("tl", "Total liabilities", hist_formula=tl, fcst=tl, bold=True)
    B.section("Equity")
    B.line("cap", "Share capital and capital surplus", hist=bs("share_capital", "capital_surplus"),
           fcst=lambda y: f"={B.ref('cap', prev(y))}")
    B.line("re", "Retained earnings", hist=bs("retained_earnings"),
           fcst=lambda y: f"={B.ref('re', prev(y))}+{I.ref('ni', y, sheet=True)}+{C.ref('div', y, sheet=True)}")
    B.line("ts", "Treasury shares", hist=bs("treasury_shares"), fcst=lambda y: f"={B.ref('ts', prev(y))}+{C.ref('bb', y, sheet=True)}")
    B.line("oce", "Other components of equity", hist=bs("other_components_equity"), fcst=lambda y: f"={B.ref('oce', prev(y))}",
           note="Held flat: no forecast of FX translation reserve")
    pe = lambda y: f"=SUM({B.ref('cap', y)}:{B.ref('oce', y)})"
    B.line("pe", "Equity attributable to owners of the parent", hist_formula=pe, fcst=pe, bold=True)
    B.line("nci", "Non-controlling interests", hist=bs("nci"), fcst=lambda y: f"={B.ref('nci', prev(y))}+{I.ref('ni_nci', y, sheet=True)}")
    B.line("te", "Total equity", hist_formula=lambda y: f"={B.ref('pe', y)}+{B.ref('nci', y)}",
           fcst=lambda y: f"={B.ref('pe', y)}+{B.ref('nci', y)}", bold=True)
    B.line("tle", "Total liabilities and equity", hist_formula=lambda y: f"={B.ref('tl', y)}+{B.ref('te', y)}",
           fcst=lambda y: f"={B.ref('tl', y)}+{B.ref('te', y)}", bold=True)
    B.line("chk", "Check: assets - liabilities - equity", hist_formula=lambda y: f"=ROUND({B.ref('ta', y)}-{B.ref('tle', y)},3)",
           fcst=lambda y: f"=ROUND({B.ref('ta', y)}-{B.ref('tle', y)},3)", note="Must be zero")
    B.section("Memo")
    nwc = lambda y: f"={B.ref('rec', y)}+{B.ref('inv', y)}+{B.ref('oca', y)}-{B.ref('pay', y)}-{B.ref('ool', y)}"
    B.line("nwc", "Operating net working capital", hist_formula=nwc, fcst=nwc)
    B.line("netcash", "Net cash (cash + deposits - borrowings - leases)",
           hist_formula=lambda y: f"={B.ref('cash', y)}+{B.ref('dep', y)}-{B.ref('borr', y)}-{B.ref('ofl', y)}",
           fcst=lambda y: f"={B.ref('cash', y)}+{B.ref('dep', y)}-{B.ref('borr', y)}-{B.ref('ofl', y)}")
    B.line("roe", "ROE (on average parent equity)",
           hist_formula=lambda y: f"={I.ref('ni', y, sheet=True)}/AVERAGE({B.ref('pe', y)},{B.ref('pe', prev(y))})" if y > HIST[0] else '=""',
           fcst=lambda y: f"={I.ref('ni', y, sheet=True)}/AVERAGE({B.ref('pe', y)},{B.ref('pe', prev(y))})", fmt="pct")

    # ---------------- Cash flow
    cf = lambda item: h(CF_, item)
    C.title("Cash-flow statement", "Indirect method. Historical condensed from data/clean/makita_cf.csv")
    C.header()
    C.section("Operating")
    C.line("profit", "Profit", hist=cf("cfo_profit"), fcst=lambda y: f"={I.ref('profit', y, sheet=True)}")
    C.line("dep", "  Depreciation ex right-of-use", fcst=lambda y: f"={sv('dep_rate')}*{B.ref('ppe', prev(y), sheet=True)}",
           note="Opening PP&E x rate")
    C.line("rou", "  Right-of-use depreciation", fcst=lambda y: f"={I.ref('rev', y, sheet=True)}*{sv('rou')}")
    C.line("da", "D&A total", hist=cf("cfo_da"), fcst=lambda y: f"={C.ref('dep', y)}+{C.ref('rou', y)}")
    for k, lab, item, bk in [("d_rec", "(Increase) / decrease in receivables", "cfo_chg_receivables", "rec"),
                             ("d_inv", "(Increase) / decrease in inventories", "cfo_chg_inventories", "inv")]:
        C.line(k, lab, hist=cf(item), fcst=lambda y, bk=bk: f"={B.ref(bk, prev(y), sheet=True)}-{B.ref(bk, y, sheet=True)}")
    C.line("d_pay", "Increase / (decrease) in payables", hist=cf("cfo_chg_payables"),
           fcst=lambda y: f"={B.ref('pay', y, sheet=True)}-{B.ref('pay', prev(y), sheet=True)}")
    C.line("d_oth", "Other working capital", fcst=lambda y: f"=({B.ref('oca', prev(y), sheet=True)}-{B.ref('oca', y, sheet=True)})"
                                                            f"+({B.ref('ool', y, sheet=True)}-{B.ref('ool', prev(y), sheet=True)})")
    C.line("cfo_oth", "Other operating items (historical: tax, interest, other)",
           hist_formula=lambda y: f"={C.ref('cfo', y)}-SUM({C.ref('profit', y)}:{C.ref('d_oth', y)})",
           note="Historical plug to reported CFO; forecast: cash taxes = tax expense")
    cfo_hist = cf("cfo")
    row = C.line("cfo", "Cash flow from operations", hist=cfo_hist,
                 fcst=lambda y: f"=SUM({C.ref('profit', y)}:{C.ref('d_oth', y)})-{C.ref('dep', y)}-{C.ref('rou', y)}", bold=True,
                 note="Forecast sums profit, D&A total and working capital (dep/ROU rows are components of D&A total)")
    C.section("Investing")
    C.line("capex", "Capex (purchase of non-current assets)", hist=cf("capex"), fcst=lambda y: f"=-{C.ref('capex_pos', y)}")
    C.line("cfi_oth", "Other investing (time deposits, investments)", hist_formula=lambda y: f"={C.ref('cfi', y)}-{C.ref('capex', y)}",
           fcst=lambda y: "=0", note="Historical plug; forecast deposits held flat")
    C.line("cfi", "Cash flow from investing", hist=cf("cfi"), fcst=lambda y: f"={C.ref('capex', y)}+{C.ref('cfi_oth', y)}", bold=True)
    C.section("Financing")
    C.line("div", "Dividends paid", hist=cf("dividends_paid"), fcst=lambda y: f"=-{I.ref('ni', y, sheet=True)}*{sv('payout')}",
           note="Forecast: payout x current-year profit (timing simplification)")
    C.line("bb", "Share buybacks", hist=cf("buybacks_net"), fcst=lambda y: f"=-{ts('buyback', y)}")
    C.line("lease", "Repayment of lease liabilities", hist=cf("lease_repayments"), fcst=lambda y: f"=-{C.ref('rou', y)}")
    C.line("d_borr", "Change in borrowings (revolver)", hist=cf("chg_st_borrowings"),
           fcst=lambda y: f"=MAX(0,{sv('min_cash')}-({C.ref('cash_beg', y)}+{C.ref('cfo', y)}+{C.ref('cfi', y)}+{C.ref('div', y)}"
                          f"+{C.ref('bb', y)}+{C.ref('lease', y)}-({B.ref('borr', prev(y), sheet=True)}-{B.ref('borr', HIST[-1], sheet=True)})))"
                          f"-({B.ref('borr', prev(y), sheet=True)}-{B.ref('borr', HIST[-1], sheet=True)})",
           note="Draws if pre-financing cash < minimum; repays prior draws first")
    C.line("cff_oth", "Other financing", hist_formula=lambda y: f"={C.ref('cff', y)}-SUM({C.ref('div', y)}:{C.ref('d_borr', y)})",
           fcst=lambda y: "=0")
    C.line("cff", "Cash flow from financing", hist=cf("cff"), fcst=lambda y: f"=SUM({C.ref('div', y)}:{C.ref('cff_oth', y)})", bold=True)
    C.section("Cash")
    C.line("fx", "FX effect on cash", hist=cf("fx_effect_cash"), fcst=lambda y: "=0", note="Not forecast")
    C.line("net", "Net change in cash", hist_formula=lambda y: f"={C.ref('cfo', y)}+{C.ref('cfi', y)}+{C.ref('cff', y)}+{C.ref('fx', y)}",
           fcst=lambda y: f"={C.ref('cfo', y)}+{C.ref('cfi', y)}+{C.ref('cff', y)}+{C.ref('fx', y)}", bold=True)
    C.line("cash_beg", "Cash at beginning of year", hist=cf("cash_begin"), fcst=lambda y: f"={B.ref('cash', prev(y), sheet=True)}")
    C.line("cash_end", "Cash at end of year", hist_formula=lambda y: f"={C.ref('cash_beg', y)}+{C.ref('net', y)}",
           fcst=lambda y: f"={C.ref('cash_beg', y)}+{C.ref('net', y)}", bold=True)
    C.section("Memo")
    C.line("capex_pos", "Capex (positive)", hist_formula=lambda y: f"=-{C.ref('capex', y)}",
           fcst=lambda y: f"={I.ref('rev', y, sheet=True)}*{ts('capex', y)}")
    C.line("fcf", "Free cash flow (CFO + capex + lease repayments)",
           hist_formula=lambda y: f"={C.ref('cfo', y)}+{C.ref('capex', y)}+{C.ref('lease', y)}",
           fcst=lambda y: f"={C.ref('cfo', y)}+{C.ref('capex', y)}+{C.ref('lease', y)}", bold=True)

    # ---------------- DCF
    V.title("DCF: unlevered free cash flow", "Cash flows from 1 Jul 2026 (after the 30 Jun 2026 balance sheet), mid-period discounting to the valuation date")
    V.header()
    V.section("Unlevered free cash flow (lease payments deducted, so leases are excluded from the equity bridge)")
    V.line("ebit", "EBIT (operating profit)", fcst=lambda y: f"={I.ref('op', y, sheet=True)}")
    V.line("tax", "Taxes on EBIT", fcst=lambda y: f"=-{V.ref('ebit', y)}*{sv('tax')}")
    V.line("nopat", "NOPAT", fcst=lambda y: f"={V.ref('ebit', y)}+{V.ref('tax', y)}", bold=True)
    V.line("dep", "+ Depreciation ex right-of-use", fcst=lambda y: f"={C.ref('dep', y, sheet=True)}",
           note="ROU depreciation is not added back: it stands in for lease cost")
    V.line("capex", "- Capex", fcst=lambda y: f"={C.ref('capex', y, sheet=True)}")
    V.line("dnwc", "- Increase in net working capital", fcst=lambda y: f"={B.ref('nwc', prev(y), sheet=True)}-{B.ref('nwc', y, sheet=True)}")
    V.line("ufcf", "Unlevered free cash flow", fcst=lambda y: f"={V.ref('nopat', y)}+{V.ref('dep', y)}+{V.ref('capex', y)}+{V.ref('dnwc', y)}", bold=True)
    V.line("reinv", "  Reinvestment (dep - capex - dNWC) % revenue", fmt="pct",
           fcst=lambda y: f"=({V.ref('dep', y)}+{V.ref('capex', y)}+{V.ref('dnwc', y)})/{I.ref('rev', y, sheet=True)}")
    V.section("Timing")
    V.line("start", "Period start", fmt="date",
           fcst=lambda y: f"={MK['bs_date']}+1" if y == FCST[0] else f"=DATE({y - 1},4,1)")
    V.line("end", "Period end", fmt="date", fcst=lambda y: f"=DATE({y},3,31)")
    V.line("frac", "Fraction of fiscal year included", fmt="dec3",
           fcst=lambda y: f"=({V.ref('end', y)}-{V.ref('start', y)}+1)/(DATE({y},3,31)-DATE({y - 1},3,31))")
    V.line("t", "Years from valuation date to mid-period", fmt="dec3",
           fcst=lambda y: f"=(({V.ref('start', y)}+{V.ref('end', y)})/2-{MK['val_date']})/365.25")
    V.line("fcf_adj", "UFCF included", fcst=lambda y: f"={V.ref('ufcf', y)}*{V.ref('frac', y)}")
    V.line("df", "Discount factor", fmt="dec3", fcst=lambda y: f"=1/(1+{MK['wacc']})^{V.ref('t', y)}")
    V.line("pv", "Present value", fcst=lambda y: f"={V.ref('fcf_adj', y)}*{V.ref('df', y)}", bold=True)
    V.line("n", "Years of growth from FY2026 (for sensitivity)", fmt="dec", fcst=lambda y: f"={y - HIST[-1]}")

    rng = lambda key: f"{V.ref(key, FCST[0], absolute=True)}:{V.ref(key, FCST[-1], absolute=True)}"
    lastc = lambda key: V.ref(key, FCST[-1], absolute=True)
    V.section("Valuation summary")
    VK = {}

    def vline(key, label, formula, fmt="num", note="", key_fmt=False):
        row = V.r
        VK[key] = f"DCF!$D${row + 1}"
        V.ws.write(row, 0, label, f["bold"] if key_fmt else f["label"])
        V.ws.write_formula(row, 3, formula, f[fmt + ("_key" if key_fmt else "")])
        if note:
            V.ws.write(row, NOTE_COL, note, f["note"])
        V.r += 1

    vline("wacc", "WACC", f"={MK['wacc']}", "pct", "Assumptions sheet")
    vline("g", "Terminal growth", f"={sv('tg')}", "pct")
    vline("mult", "Exit EV/EBITDA", f"={sv('exit')}", "x")
    vline("t_end", "Years to end of FY2031", f"=({lastc('end')}-{MK['val_date']})/365.25", "dec3")
    vline("sum_pv", "PV of explicit-period UFCF", f"=SUM({rng('pv')})")
    vline("bridge", "+ Cash + deposits + investments - borrowings - NCI - declared interim dividend",
          f"={MK['cash_q']}+{MK['dep_q']}+{MK['inv_q']}-{MK['borr_q']}-{MK['nci_q']}-{MK['interim']}*{MK['shares']}",
          note="30 Jun 2026 balance sheet; investments at book; the stock is already ex the JPY 79 interim")
    V.r += 1
    V.ws.write(V.r, 0, "Gordon growth method", f["sec"]); V.r += 1
    vline("ronic", "RONIC (terminal)", f"={sv('ronic')}", "pct")
    vline("fcf_t1", "Terminal-year FCF: NOPAT FY2031 x (1 + g) x (1 - g / RONIC)",
          f"={lastc('nopat')}*(1+{VK['g']})*(1-{VK['g']}/{VK['ronic']})",
          note="Value-driver formula (McKinsey): avoids assuming FY2031's reinvestment rate forever")
    vline("tv_g", "Terminal value at end FY2031", f"={VK['fcf_t1']}/({VK['wacc']}-{VK['g']})")
    vline("pv_tv_g", "PV of terminal value", f"={VK['tv_g']}/(1+{VK['wacc']})^{VK['t_end']}")
    vline("ev_g", "Enterprise value", f"={VK['sum_pv']}+{VK['pv_tv_g']}")
    vline("eq_g", "Equity value", f"={VK['ev_g']}+{VK['bridge']}")
    vline("ps_g", "Value per share (JPY)", f"={VK['eq_g']}/{MK['shares']}", "yen", key_fmt=True)
    vline("tv_share_g", "Terminal value share of EV", f"={VK['pv_tv_g']}/{VK['ev_g']}", "pct")
    vline("impl_mult", "Implied exit EV/EBITDA (TV + leases) / EBITDA FY2031",
          f"=({VK['tv_g']}+{MK['lease_fy']})/{I.ref('ebitda', FCST[-1], absolute=True, sheet=True)}", "x",
          "Cross-check against the exit-multiple assumption")
    V.r += 1
    V.ws.write(V.r, 0, "Exit-multiple method", f["sec"]); V.r += 1
    vline("tv_x", "Terminal value at end FY2031 (EBITDA x multiple - leases)",
          f"={I.ref('ebitda', FCST[-1], absolute=True, sheet=True)}*{VK['mult']}-{MK['lease_fy']}",
          note="Multiple applies to IFRS 16 EBITDA, which gives EV incl. leases; leases deducted for consistency")
    vline("pv_tv_x", "PV of terminal value", f"={VK['tv_x']}/(1+{VK['wacc']})^{VK['t_end']}")
    vline("ev_x", "Enterprise value", f"={VK['sum_pv']}+{VK['pv_tv_x']}")
    vline("eq_x", "Equity value", f"={VK['ev_x']}+{VK['bridge']}")
    vline("ps_x", "Value per share (JPY)", f"={VK['eq_x']}/{MK['shares']}", "yen", key_fmt=True)
    vline("impl_g", "Implied perpetual growth (value-driver, solves TV = NOPAT(1+g)(1-g/RONIC)/(WACC-g))",
          (lambda N, R, TV, W: f"=(-(({N}/{R})-{N}-{TV})-SQRT(((({N}/{R})-{N}-{TV}))^2-4*({N}/{R})*({TV}*{W}-{N})))/(2*{N}/{R})")(
              lastc('nopat'), VK['ronic'], VK['tv_x'], VK['wacc']), "pct", "Cross-check against the Gordon growth assumption")
    V.r += 1
    V.ws.write(V.r, 0, "Target price and rating", f["sec"]); V.r += 1
    vline("fv", "Fair value today: average of the two DCF methods", f"=AVERAGE({VK['ps_g']},{VK['ps_x']})", "yen")
    vline("ke", "Cost of equity", f"={MK['ke']}", "pct")
    vline("dps12", "Dividends expected in next 12 months (JPY)",
          f"=({I.ref('dps', FCST[0], absolute=True, sheet=True)}-{MK['interim']})+0.5*{I.ref('dps', FCST[1], absolute=True, sheet=True)}", "dec",
          "FY2027 year-end (paid Jun 2027) + FY2028 interim (record Sep 2027)")
    vline("tp_raw", "12-month value: fair value x (1 + Ke) - dividends", f"={VK['fv']}*(1+{VK['ke']})-{VK['dps12']}", "yen")
    vline("tp", "12-month target price (rounded to JPY 50)", f"=MROUND({VK['tp_raw']},50)", "yen", key_fmt=True)
    vline("price", "Current price", f"={MK['price']}", "yen")
    vline("upside", "Upside / (downside)", f"={VK['tp']}/{VK['price']}-1", "pct", key_fmt=True)
    vline("tsr", "Expected total return incl. dividends", f"=({VK['tp']}+{VK['dps12']})/{VK['price']}-1", "pct")
    V.ws.write(V.r, 0, "Rating", f["bold"])
    VK["rating"] = f"DCF!$D${V.r + 1}"
    V.ws.write_formula(V.r, 3, f'=IF({VK["upside"]}>0.15,"BUY",IF({VK["upside"]}<-0.1,"SELL","HOLD"))', f["switch"])
    V.ws.write(V.r, NOTE_COL, "Rule: BUY if upside > +15%; SELL if upside < -10%; otherwise HOLD", f["note"])
    V.r += 1

    # ---------------- Comps
    P.title("Trading comparables", "Last reported fiscal year; all values converted to JPY m at 29 Oct 2026 FX (peers_financials.csv)")
    P.ws.set_column(1, 1, 24)
    cols = ["Ticker", "Company", "FY end", "Mkt cap", "EV", "Revenue", "EBITDA", "Net income", "Equity",
            "EV/Sales", "EV/EBITDA", "P/E", "P/B", "EBITDA mgn"]
    P.r = 3
    for j, c in enumerate(cols):
        P.ws.write(P.r, j, c, f["hdr_r"] if j > 2 else f["hdr"])
    P.r += 1
    peers = d["peers"][d["peers"].ticker != "6586.T"]
    first = P.r
    for _, p in peers.iterrows():
        r = P.r
        P.ws.write(r, 0, p.ticker, f["text_in"]); P.ws.write(r, 1, p["name"], f["text_in"]); P.ws.write(r, 2, p.fy_end, f["text_in"])
        for j, k in enumerate(["market_cap_jpy_m", "ev_jpy_m", "revenue_jpy_m", "ebitda_jpy_m", "net_income_jpy_m", "equity_jpy_m"]):
            P.ws.write_number(r, 3 + j, float(p[k]), f["num_in"])
        rr = r + 1
        P.ws.write_formula(r, 9, f"=E{rr}/F{rr}", f["x"]); P.ws.write_formula(r, 10, f"=E{rr}/G{rr}", f["x"])
        P.ws.write_formula(r, 11, f"=D{rr}/H{rr}", f["x"]); P.ws.write_formula(r, 12, f"=D{rr}/I{rr}", f["x"])
        P.ws.write_formula(r, 13, f"=G{rr}/F{rr}", f["pct"])
        P.r += 1
    last = P.r - 1
    PK = {}
    for lab, fn in [("Median", "MEDIAN"), ("Mean", "AVERAGE"), ("25th percentile", "Q1"), ("75th percentile", "Q3")]:
        r = P.r
        P.ws.write(r, 1, lab, f["bold"])
        for j in range(9, 14):
            colL = chr(ord("A") + j)
            rngp = f"{colL}{first + 1}:{colL}{last + 1}"
            frm = {"Q1": f"=QUARTILE({rngp},1)", "Q3": f"=QUARTILE({rngp},3)"}.get(fn, f"={fn}({rngp})")
            P.ws.write_formula(r, j, frm, f["b_pct" if j == 13 else "b_x"])
            PK[(lab, cols[j])] = f"Comps!${colL}${r + 1}"
        P.r += 1
    P.r += 1
    P.ws.write(P.r, 0, "Makita (live, LTM to 30 Jun 2026)", f["sec"]); P.r += 1
    MKC = {}

    def pl(key, label, formula, fmt="num", note=""):
        r = P.r
        MKC[key] = f"Comps!$D${r + 1}"
        P.ws.write(r, 1, label, f["label"])
        if isinstance(formula, str):
            P.ws.write_formula(r, 3, formula, f[fmt])
        else:
            P.ws.write_number(r, 3, float(formula), f[fmt + "_in"])
        if note:
            P.ws.write(r, 5, note, f["note"])
        P.r += 1

    pl("mcap", "Market cap", f"={MK['mcap']}")
    pl("ev", "EV (incl. leases; cash + deposits deducted)", f"={MK['mcap']}+{MK['debt_mv']}-{MK['cash_q']}-{MK['dep_q']}+{MK['nci_q']}",
       note="Same EV definition as peers (yfinance: total debt incl. leases less cash and short-term investments)")
    pl("rev", "Revenue LTM", ltm["revenue"], note="FY2026 + 1Q FY2027 - 1Q FY2026")
    pl("ebitda", "EBITDA LTM (IFRS 16)", ltm["operating_profit"] + ltm["cfo_da"])
    pl("ni", "Net income LTM", ltm["ni_parent"])
    pl("eq", "Parent equity 30 Jun 2026", bsq["equity_parent"])
    pl("ev_ebitda", "EV/EBITDA", f"={MKC['ev']}/{MKC['ebitda']}", "x")
    pl("pe", "P/E", f"={MKC['mcap']}/{MKC['ni']}", "x")
    pl("ev_sales", "EV/Sales", f"={MKC['ev']}/{MKC['rev']}", "x")
    pl("pb", "P/B", f"={MKC['mcap']}/{MKC['eq']}", "x")
    hm = d["hm"]
    hm_ex = hm[hm.fy != "FY2023"]
    pl("own_med", "Makita own EV/EBITDA, FY2019-26 median", float(hm.ev_ebitda.median()), "x", note="makita_hist_multiples.csv")
    pl("own_lo", "Makita own EV/EBITDA, FY2019-26 low (ex FY2023 trough)", float(hm_ex.ev_ebitda.min()), "x")
    pl("own_hi", "Makita own EV/EBITDA, FY2019-26 high (ex FY2023 trough)", float(hm_ex.ev_ebitda.max()), "x")
    P.r += 1
    P.ws.write(P.r, 0, "Implied Makita value per share (JPY)", f["sec"]); P.r += 1
    ev2ps = lambda m: f"=({m}*{MKC['ebitda']}-{MKC['ev']}+{MKC['mcap']})/{MK['shares']}"
    for lab, q in [("Median", "med"), ("25th percentile", "lo"), ("75th percentile", "hi")]:
        pl(f"ps_evebitda_{q}", f"EV/EBITDA at peer {lab.lower()}", ev2ps(PK[(lab, 'EV/EBITDA')]), "yen")
        pl(f"ps_pe_{q}", f"P/E at peer {lab.lower()}", f"={PK[(lab, 'P/E')]}*{MKC['ni']}/{MK['shares']}", "yen")
    pl("ps_own_lo", "EV/EBITDA at own FY2019-26 low", ev2ps(MKC["own_lo"]), "yen")
    pl("ps_own_hi", "EV/EBITDA at own FY2019-26 high", ev2ps(MKC["own_hi"]), "yen")
    pl("ps_own_med", "EV/EBITDA at own FY2019-26 median", ev2ps(MKC["own_med"]), "yen")
    P.ws.write(P.r + 1, 1, "Peer set: TTI, Stanley Black & Decker, Chervon (cordless power tools and OPE); Husqvarna, Toro, Yamabiko "
                           "(OPE); Max (Japanese tools). Bosch and Hilti are private. SWK and Toro verified vs SEC 10-K.", f["note"])

    # ---------------- Sensitivity
    S.title("Sensitivity tables (live formulas: each cell recomputes the DCF)", "Centre cells equal the DCF sheet; see Checks")
    S.r = 3
    wacc_steps = [-0.01, -0.005, 0, 0.005, 0.01]
    g_steps = [-0.01, -0.005, 0, 0.005, 0.01]
    SK = {}

    def grid(title, row_label, col_label, row_vals, col_vals, row_fmt, col_fmt, cell, key):
        S.ws.write(S.r, 0, title, f["sec"]); S.r += 1
        S.ws.write(S.r, 1, f"{row_label} \\ {col_label}", f["unit"])
        hr = S.r
        for j, cv in enumerate(col_vals):
            S.ws.write_formula(hr, 2 + j, cv, f[col_fmt + "_key"])
        for i, rv in enumerate(row_vals):
            r = hr + 1 + i
            S.ws.write_formula(r, 1, rv, f[row_fmt + "_key"])
            for j in range(len(col_vals)):
                S.ws.write_formula(r, 2 + j, cell(xl_rowcol_to_cell(r, 1, col_abs=True), xl_rowcol_to_cell(hr, 2 + j, row_abs=True)), f["yen"])
        SK[key] = (hr, len(row_vals), len(col_vals))
        S.r = hr + len(row_vals) + 2

    pv_at = lambda w: f"SUMPRODUCT(DCF!{rng('fcf_adj')}/(1+{w})^DCF!{rng('t')})"
    tend = VK["t_end"]
    grid("1. Value per share (Gordon): WACC vs terminal growth", "WACC", "g",
         [f"={VK['wacc']}+({s})" for s in wacc_steps], [f"={VK['g']}+({s})" for s in g_steps], "pct", "pct",
         lambda w, g: f"=({pv_at(w)}+DCF!{lastc('nopat')}*(1+{g})*(1-{g}/{VK['ronic']})/({w}-{g})/(1+{w})^{tend}+{VK['bridge']})/{MK['shares']}",
         "wacc_g")
    grid("2. Value per share (exit multiple): WACC vs exit EV/EBITDA", "WACC", "Multiple",
         [f"={VK['wacc']}+({s})" for s in wacc_steps], [f"={VK['mult']}+({s})" for s in [-2, -1, 0, 1, 2]], "pct", "x",
         lambda w, m: f"=({pv_at(w)}+({I.ref('ebitda', FCST[-1], absolute=True, sheet=True)}*{m}-{MK['lease_fy']})/(1+{w})^{tend}+{VK['bridge']})/{MK['shares']}",
         "wacc_x")
    # 3. operating: uniform local-currency growth uplift vs steady-state EBIT margin (Gordon, base WACC and g)
    revr, taxr = f"IS!{I.ref('rev', FCST[0], absolute=True)}:{I.ref('rev', FCST[-1], absolute=True)}", sv("tax")
    fcf_alt = lambda dg, m: (f"IS!{I.ref('rev', FCST[0], absolute=True)}:{I.ref('rev', FCST[-1], absolute=True)}*(1+{dg})^DCF!{rng('n')}"
                             f"*({m}*(1-{taxr})+DCF!{rng('reinv')})")
    last_alt = lambda dg, m: (f"IS!{I.ref('rev', FCST[-1], absolute=True)}*(1+{dg})^DCF!{lastc('n')}*{m}*(1-{taxr})"
                              f"*(1-{g0}/{VK['ronic']})")  # terminal NOPAT x (1 - g/RONIC); (1+g) applied below
    w0, g0 = VK["wacc"], VK["g"]
    grid("3. Value per share (Gordon): annual revenue growth uplift vs steady-state operating margin", "Growth +/-", "OP margin",
         [f"={s}" for s in [-0.02, -0.01, 0, 0.01, 0.02]], [f"={m}" for m in [0.11, 0.125, 0.14, 0.155, 0.17]], "pct", "pct",
         lambda dg, m: (f"=(SUMPRODUCT({fcf_alt(dg, m)}*DCF!{rng('frac')}/(1+{w0})^DCF!{rng('t')})"
                        f"+{last_alt(dg, m)}*(1+{g0})/({w0}-{g0})/(1+{w0})^{tend}+{VK['bridge']})/{MK['shares']}"),
         "growth_margin")
    S.ws.write(S.r, 0, "Table 3 holds the operating margin constant in all forecast years and scales revenue by (1 + uplift)^n; "
                       "reinvestment stays at the model's % of revenue. It is a linearised view, so its centre differs from the full model.",
               f["note"])
    S.r += 2

    # FX sensitivity: EPS FY2028E under +/-10% yen move is computed in Python by re-running scenarios (see model_outputs.json)

    # ---------------- Football field
    FF.title("Football field (JPY per share)", "Ranges feed the chart in report/charts/07_football_field.png")
    FF.r = 3
    for j, c in enumerate(["Method", "", "Low", "High", "Basis"]):
        FF.ws.write(FF.r, j, c, f["hdr"])
    FF.r += 1
    FK = {}
    g1, g2 = SK["wacc_g"], SK["wacc_x"]
    cell_at = lambda gr, i, j: f"Sensitivity!{xl_rowcol_to_cell(gr[0] + 1 + i, 2 + j, row_abs=True, col_abs=True)}"
    rows_ff = [
        ("dcf_g", "DCF, Gordon growth", f"=MIN({cell_at(g1, 1, 1)}:{cell_at(g1, 3, 3)})", f"=MAX({cell_at(g1, 1, 1)}:{cell_at(g1, 3, 3)})",
         "WACC +/-0.5pp, g +/-0.5pp"),
        ("dcf_x", "DCF, exit multiple", f"=MIN({cell_at(g2, 1, 1)}:{cell_at(g2, 3, 3)})", f"=MAX({cell_at(g2, 1, 1)}:{cell_at(g2, 3, 3)})",
         "WACC +/-0.5pp, multiple +/-1x"),
        ("cmp_ev", "Peers EV/EBITDA (LTM)", f"={MKC['ps_evebitda_lo']}", f"={MKC['ps_evebitda_hi']}", "Peer 25th-75th percentile"),
        ("cmp_pe", "Peers P/E (LTM)", f"={MKC['ps_pe_lo']}", f"={MKC['ps_pe_hi']}", "Peer 25th-75th percentile"),
        ("own", "Makita own EV/EBITDA history", f"={MKC['ps_own_lo']}", f"={MKC['ps_own_hi']}", "FY2019-26 range ex FY2023"),
        ("w52", "52-week trading range", float(snap["low_52w_close"]), float(snap["high_52w_close"]), "Closing prices"),
    ]
    for key, lab, lo, hi, basis in rows_ff:
        r = FF.r
        FK[key] = r
        FF.ws.write(r, 0, lab, f["label"])
        for j, v in [(2, lo), (3, hi)]:
            if isinstance(v, str):
                FF.ws.write_formula(r, j, v, f["yen"])
            else:
                FF.ws.write_number(r, j, v, f["yen_in"])
        FF.ws.write(r, 4, basis, f["note"])
        FF.r += 1
    FF.r += 1
    FF.ws.write(FF.r, 0, "Current price", f["bold"]); FF.ws.write_formula(FF.r, 2, f"={MK['price']}", f["yen"]); FF.r += 1
    FF.ws.write(FF.r, 0, "12-month target price", f["bold"]); FF.ws.write_formula(FF.r, 2, f"={VK['tp']}", f["yen_key"]); FF.r += 1

    # ---------------- Checks
    K.title("Model checks", "Every check must read TRUE. Forecast checks exact; historical checks allow JPY 2m source rounding.")
    K.header("Check", years=True)
    checks = []

    def kline(key, label, hist_f=None, fcst_f=None, single=None, note=""):
        r = K.r
        K.ws.write(r, 0, label, f["label"])
        if single is not None:
            K.ws.write_formula(r, 2, single, f["ok"])
            checks.append(xl_rowcol_to_cell(r, 2))
        for y in (HIST if hist_f else []):
            K.ws.write_formula(r, COL[y], hist_f(y), f["ok"]); checks.append(xl_rowcol_to_cell(r, COL[y]))
        for y in (FCST if fcst_f else []):
            K.ws.write_formula(r, COL[y], fcst_f(y), f["ok"]); checks.append(xl_rowcol_to_cell(r, COL[y]))
        if note:
            K.ws.write(r, NOTE_COL, note, f["note"])
        K.r += 1

    ab = lambda x: f"=ABS({x})<0.5"      # forecast: exact (formula-driven)
    abh = lambda x: f"=ABS({x})<=2"      # historical: company rounds each reported line to JPY 1m
    kline("bs", "Balance sheet balances (forecast: imbalance unchanged from FY2026A rounding)",
          lambda y: abh(B.ref("chk", y, sheet=True)),
          lambda y: ab(f"{B.ref('chk', y, sheet=True)}-{B.ref('chk', HIST[-1], sheet=True)}"),
          note="Historical totals are re-added from reported lines, so JPY 1-2m rounding differences can appear")
    kline("cash_tie", "CF ending cash = BS cash", lambda y: abh(f"{C.ref('cash_end', y, sheet=True)}-{B.ref('cash', y, sheet=True)}"),
          lambda y: ab(f"{C.ref('cash_end', y, sheet=True)}-{B.ref('cash', y, sheet=True)}"))
    kline("cash_roll", "CF beginning cash = prior BS cash", None,
          lambda y: ab(f"{C.ref('cash_beg', y, sheet=True)}-{B.ref('cash', prev(y), sheet=True)}"))
    kline("rev_tie", "Revenue build = IS revenue", lambda y: abh(f"{D.ref('rev', y, sheet=True)}-{I.ref('rev', y, sheet=True)}"),
          lambda y: ab(f"{D.ref('rev', y, sheet=True)}-{I.ref('rev', y, sheet=True)}"), note="Historical: regions sum to reported revenue")
    kline("ni_tie", "IS profit = CF profit", lambda y: abh(f"{I.ref('profit', y, sheet=True)}-{C.ref('profit', y, sheet=True)}"),
          lambda y: ab(f"{I.ref('profit', y, sheet=True)}-{C.ref('profit', y, sheet=True)}"))
    kline("re_roll", "Retained earnings roll-forward (forecast)", None,
          lambda y: ab(f"{B.ref('re', y, sheet=True)}-{B.ref('re', prev(y), sheet=True)}-{I.ref('ni', y, sheet=True)}-{C.ref('div', y, sheet=True)}"))
    kline("dcf_ebit", "DCF EBIT = IS operating profit", None, lambda y: ab(f"{V.ref('ebit', y, sheet=True)}-{I.ref('op', y, sheet=True)}"))
    kline("dcf_capex", "DCF capex = CF capex", None, lambda y: ab(f"{V.ref('capex', y, sheet=True)}-{C.ref('capex', y, sheet=True)}"))
    kline("dcf_nwc", "DCF change in NWC = CF working-capital lines", None,
          lambda y: ab(f"{V.ref('dnwc', y, sheet=True)}-({C.ref('d_rec', y, sheet=True)}+{C.ref('d_inv', y, sheet=True)}"
                       f"+{C.ref('d_pay', y, sheet=True)}+{C.ref('d_oth', y, sheet=True)})"))
    kline("dcf_tax", "DCF tax rate = model tax rate", None, lambda y: f"=ABS(-{V.ref('tax', y, sheet=True)}/{V.ref('ebit', y, sheet=True)}-{sv('tax')})<0.0001")
    kline("lease", "Lease repayments = ROU depreciation (ROU asset flat)", None,
          lambda y: ab(f"{C.ref('lease', y, sheet=True)}+{C.ref('rou', y, sheet=True)}"))
    kline("wacc_w", "WACC weights sum to 100%", single=f"=ABS({MK['we']}+{MK['wd']}-1)<0.000001")
    kline("g_lt_w", "Terminal growth < WACC", single=f"={VK['g']}<{VK['wacc']}")
    kline("g_lt_rf", "Terminal growth <= risk-free rate", single=f"={VK['g']}<={MK['rf']}")
    kline("impl_g", "Exit-multiple implied growth <= risk-free rate (TV methods consistent)",
          single=f"={VK['impl_g']}<={MK['rf']}")
    hr, nr, nc = SK["wacc_g"]
    kline("sens_g", "Sensitivity 1 centre = DCF Gordon value",
          single=f"=ABS(Sensitivity!{xl_rowcol_to_cell(hr + 1 + nr // 2, 2 + nc // 2)}-{VK['ps_g']})<0.01")
    hr, nr, nc = SK["wacc_x"]
    kline("sens_x", "Sensitivity 2 centre = DCF exit value",
          single=f"=ABS(Sensitivity!{xl_rowcol_to_cell(hr + 1 + nr // 2, 2 + nc // 2)}-{VK['ps_x']})<0.01")
    kline("scen", "Scenario switch is 1, 2 or 3", single="=OR(Scenario=1,Scenario=2,Scenario=3)")
    kline("frac", "Stub fraction between 0 and 1", single=f"=AND(DCF!{V.ref('frac', FCST[0])}>0,DCF!{V.ref('frac', FCST[0])}<=1)")
    K.r += 1
    K.ws.write(K.r, 0, "MASTER CHECK", f["bold"])
    KK = f"Checks!$C${K.r + 1}"
    K.ws.write_formula(K.r, 2, f'=IF(AND({",".join(checks)}),"ALL CHECKS PASS","CHECK FAILURES")', f["switch"])
    K.ws.write(K.r, 4, f"{len(checks)} individual checks", f["note"])

    # ---------------- Cover
    cover.set_column(0, 0, 4); cover.set_column(1, 1, 44); cover.set_column(2, 2, 18); cover.set_column(3, 3, 80)
    cover.write(1, 1, "Makita Corporation (6586 JP): initiation model", f["title"])
    cover.write(2, 1, "Student research project, not investment advice. IFRS, JPY m, FY ends 31 March.", f["sub"])
    items = [("Scenario", "=Assumptions!E4", "ok"), ("Rating", f"={VK['rating']}", "ok"),
             ("12-month target price (JPY)", f"={VK['tp']}", "yen_key"), ("Current price (JPY)", f"={MK['price']}", "yen"),
             ("Upside / (downside)", f"={VK['upside']}", "pct"), ("DCF value, Gordon (JPY)", f"={VK['ps_g']}", "yen"),
             ("DCF value, exit multiple (JPY)", f"={VK['ps_x']}", "yen"), ("WACC", f"={MK['wacc']}", "pct"),
             ("Cost of equity", f"={MK['ke']}", "pct"), ("Model checks", f"={KK}", "ok")]
    for i, (lab, frm, fm) in enumerate(items):
        cover.write(4 + i, 1, lab, f["bold"]); cover.write_formula(4 + i, 2, frm, f[fm])
    notes = ["How to use: change Assumptions!D4 (1 Base, 2 Bull, 3 Bear). Everything recalculates.",
             "Colour code: blue = input or historical fact; black = formula; tinted = key output.",
             "Sheets: Assumptions > Drivers (regional revenue x FX) > IS > BS > CF > DCF > Comps > Sensitivity > Football > Checks.",
             "Sources: Makita earnings releases and audited statements; yfinance; MoF JGB; Damodaran. See SOURCES.md.",
             "Built by src/build_model.py and recalculated in LibreOffice; values are cached so the file opens with results."]
    for i, n in enumerate(notes):
        cover.write(16 + i, 1, n, f["note"])
    flush()
    wb.close()
    return dict(VK=VK, MK=MK, KK=KK, SK=SK, FK=FK, I=I.rows, B=B.rows, C=C.rows, V=V.rows, D=D.rows, MKC=MKC, PK=PK)


# ----------------------------------------------------------------------------------------------- read back
def cellv(wb, ref):
    sheet, cell = ref.replace("$", "").split("!")
    return wb[sheet.strip("'")][cell].value


def read_outputs(path: Path, refs):
    wb = openpyxl.load_workbook(path, data_only=True)
    VK, MK = refs["VK"], refs["MK"]
    out = {k: cellv(wb, v) for k, v in VK.items()}
    out.update({"mk_" + k: cellv(wb, v) for k, v in MK.items() if k not in ("val_date", "bs_date")})
    out["checks"] = cellv(wb, refs["KK"])
    out["comps"] = {k: cellv(wb, v) for k, v in refs["MKC"].items()}
    out["peer_stats"] = {f"{a}|{b}": cellv(wb, v) for (a, b), v in refs["PK"].items()}
    line = lambda sheet, rows, key: {f"FY{y}": wb[sheet].cell(rows[key] + 1, COL[y] + 1).value for y in YEARS}
    out["is"] = {k: line("IS", refs["I"], k) for k in ["rev", "rev_g", "gp", "gm", "op", "opm", "ni", "eps", "dps", "ebitda", "ebitda_m", "da"]}
    out["bs"] = {k: line("BS", refs["B"], k) for k in ["cash", "inv", "nwc", "netcash", "pe", "roe", "chk"]}
    out["cf"] = {k: line("CF", refs["C"], k) for k in ["cfo", "capex", "fcf", "div", "bb"]}
    out["dcf"] = {k: {f"FY{y}": wb["DCF"].cell(refs["V"][k] + 1, COL[y] + 1).value for y in FCST}
                  for k in ["ebit", "nopat", "dep", "capex", "dnwc", "ufcf", "frac", "t", "df", "pv"]}
    out["drivers"] = {k: line("Drivers", refs["D"], k) for k in ["eu", "jp", "na", "rw", "rev", "eu_lc", "jp_lc", "na_lc", "rw_lc", "usd", "eur"]}
    grids = {}
    ws = wb["Sensitivity"]
    for key, (hr, nr, nc) in refs["SK"].items():
        grids[key] = {"cols": [ws.cell(hr + 1, 3 + j).value for j in range(nc)],
                      "rows": [ws.cell(hr + 2 + i, 2).value for i in range(nr)],
                      "values": [[ws.cell(hr + 2 + i, 3 + j).value for j in range(nc)] for i in range(nr)]}
    out["sensitivity"] = grids
    ws = wb["Football"]
    out["football"] = {k: {"label": ws.cell(r + 1, 1).value, "low": ws.cell(r + 1, 3).value, "high": ws.cell(r + 1, 4).value,
                           "basis": ws.cell(r + 1, 5).value} for k, r in refs["FK"].items()}
    # list failing checks, if any
    ws = wb["Checks"]
    fails = []
    for row in ws.iter_rows(min_row=5):
        for c in row[2:12]:
            if c.value is False:
                fails.append(f"{row[0].value} [{ws.cell(4, c.column).value or ''}]")
    out["check_failures"] = fails
    return out


def main():
    MODEL.mkdir(exist_ok=True)
    d = load_hist()
    tmp = MODEL / "_build"
    tmp.mkdir(exist_ok=True)
    results = {}
    for k, name in [(2, "Bull"), (3, "Bear"), (1, "Base")]:
        raw = tmp / f"model_{name}.xlsx"
        refs = build(d, k, raw)
        calc = tmp / f"model_{name}_calc.xlsx"
        recalc(raw, calc)
        results[name] = read_outputs(calc, refs)
        r = results[name]
        print(f"{name:5s} checks={r['checks']} Gordon={r['ps_g']:,.0f} exit={r['ps_x']:,.0f} TP={r['tp']:,.0f} "
              f"upside={r['upside']:+.1%} rating={r['rating']} WACC={r['wacc']:.2%}")
        if r["check_failures"]:
            print("   FAILED:", r["check_failures"][:20])
    shutil.copy(tmp / "model_Base_calc.xlsx", OUT)
    (CLEAN / "model_outputs.json").write_text(json.dumps(results, indent=1, default=str))
    shutil.rmtree(tmp)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
